"""Validate the final Office/PDF/video handoff and record exact artifact hashes."""

import hashlib
import json
import subprocess
import zipfile
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from docx import Document
from PIL import Image
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[2]
REPORT = ROOT / "docs/report"
OUT = REPORT / "generated"


def command(*args):
    return subprocess.check_output(args, text=True, timeout=60)


doc = Document(OUT / "oldshop_report.docx")
headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
for i in range(1, 10):
    assert any(h.startswith(str(i) + ". ") for h in headings), i
for appendix in "ABC":
    assert any(h.startswith("Phụ lục " + appendix + ".") for h in headings), appendix
assert "<!--" not in "\n".join(p.text for p in doc.paragraphs)

slides = Presentation(OUT / "oldshop_slides.pptx")
assert len(slides.slides) == 18
assert all(
    s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip()
    for s in slides.slides
)
pdf_pages = {}
for name in ("oldshop_report.pdf", "oldshop_slides.pdf"):
    info = command("pdfinfo", str(OUT / name))
    pages = next(
        int(line.split(":", 1)[1].strip())
        for line in info.splitlines()
        if line.startswith("Pages:")
    )
    assert pages == 18 if "slides" in name else pages >= 12
    pdf_pages[name] = pages
    text = command("pdftotext", "-layout", str(OUT / name), "-")
    assert "\ufffd" not in text and "<!--" not in text

video = json.loads(
    command(
        "ffprobe",
        "-v",
        "quiet",
        "-show_format",
        "-show_streams",
        "-of",
        "json",
        str(OUT / "oldshop_demo.mp4"),
    )
)
duration = float(video["format"]["duration"])
assert 300 <= duration <= 420
assert (video["streams"][0]["width"], video["streams"][0]["height"]) == (1280, 720)
rehearsal = json.loads((ROOT / "docs/evidence/day7_rehearsal.json").read_text())
assert rehearsal["result"] == "PASS" and len(rehearsal["stages"]) == 9
assert abs(rehearsal["video_seconds"] - duration) < 0.1
for name in [
    "01_stack",
    "02_postgres",
    "03_mongo",
    "04_product",
    "05_minio",
    "06_etl",
    "07_explain",
    "08_node_down",
    "09_recovered",
]:
    with Image.open(REPORT / "screenshots" / (name + ".png")) as image:
        assert image.size == (1280, 720)

artifacts = [
    "oldshop_report.docx",
    "oldshop_report.pdf",
    "oldshop_slides.pptx",
    "oldshop_slides.pdf",
    "oldshop_demo.mp4",
    "oldshop_demo.srt",
]
hashes = {}
for name in artifacts:
    path = OUT / name
    blob = path.read_bytes()
    assert len(blob) > 100
    hashes[name] = {"bytes": len(blob), "sha256": hashlib.sha256(blob).hexdigest()}
    if name.endswith((".docx", ".pptx")):
        with zipfile.ZipFile(path) as package:
            assert package.testzip() is None

result = {
    "result": "PASS",
    "checked_at": datetime.now(ZoneInfo("Asia/Ho_Chi_Minh")).isoformat(),
    "chapters": 9,
    "appendices": 3,
    "slides": 18,
    "all_slides_have_notes": True,
    "pdf_pages": pdf_pages,
    "video_seconds": duration,
    "video_dimensions": [1280, 720],
    "rehearsal_stages": 9,
    "required_screenshots": 9,
    "artifact_hashes": hashes,
    "scope": "Artifact integrity/structure; visual layout reviewed separately, not a production certification",
}
(ROOT / "docs/evidence/day7_output_validation.json").write_text(
    json.dumps(result, ensure_ascii=False, indent=2) + "\n"
)
print(json.dumps(result, ensure_ascii=False, indent=2))

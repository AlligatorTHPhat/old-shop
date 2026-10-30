"""Build report/deck from repository evidence; generated Office/PDF artifacts stay outside Git."""

import io
import json
import os
import re
import subprocess
import tempfile
from pathlib import Path
from xml.sax.saxutils import escape

os.environ.setdefault("MPLCONFIGDIR", "/tmp/oldshop-day7-mpl")
import cairosvg
import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from markdown_it import MarkdownIt
from pptx import Presentation
from pptx.dml.color import RGBColor as PPTColor
from pptx.util import Inches
from pptx.util import Pt as PPTPt

ROOT = Path(__file__).resolve().parents[2]
REPORT = ROOT / "docs/report"
ASSETS = REPORT / "assets"
OUT = REPORT / "generated"
NAVY, TEAL, ORANGE = "14283F", "008F86", "DE763C"
LABELS = {
    "catalogue": "Catalogue",
    "search": "ILIKE search",
    "buyer_history": "Lịch sử buyer",
    "seller_orders": "Đơn seller",
    "review_history": "Review history",
    "reports": "Report rating=0",
    "cart_lookup": "Cart lookup",
    "order_pk_pruning": "Order PK",
    "log_month_pruning": "Log theo tháng",
}


def svg_diagram(name, boxes, edges, caption, width=1100, height=570):
    parts = [
        f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}">',
        '<defs><marker id="a" markerWidth="10" markerHeight="8" refX="9" refY="4" orient="auto"><path d="M0,0 L10,4 L0,8" fill="#008F86"/></marker></defs>',
        f'<rect width="{width}" height="{height}" fill="#F6F8FA"/>',
    ]
    for x1, y1, x2, y2 in edges:
        parts.append(
            f'<line x1="{x1}" y1="{y1}" x2="{x2}" y2="{y2}" stroke="#008F86" stroke-width="2.5" marker-end="url(#a)"/>'
        )
    for x, y, w, h, title, lines in boxes:
        parts += [
            f'<rect x="{x}" y="{y}" width="{w}" height="{h}" rx="10" fill="white" stroke="#C8D6DF" stroke-width="2"/>',
            f'<text x="{x + 16}" y="{y + 30}" font-family="DejaVu Sans,Noto Sans" font-size="21" font-weight="bold" fill="#{NAVY}">{escape(title)}</text>',
        ]
        for i, line in enumerate(lines):
            parts.append(
                f'<text x="{x + 16}" y="{y + 58 + 24 * i}" font-family="DejaVu Sans,Noto Sans" font-size="16" fill="#3D5266">{escape(line)}</text>'
            )
    parts.append(
        f'<text x="28" y="{height - 20}" font-family="DejaVu Sans,Noto Sans" font-size="16" fill="#546C80">{escape(caption)}</text></svg>'
    )
    (ASSETS / (name + ".svg")).write_text("\n".join(parts))


def diagrams():
    svg_diagram(
        "07_architecture",
        [
            (35, 45, 245, 110, "Angular / Nginx", ["Browser → /api + /media"]),
            (
                370,
                45,
                310,
                110,
                "FastAPI",
                ["Core PG · metadata Mongo", "Gateway ảnh · readiness"],
            ),
            (
                35,
                230,
                290,
                125,
                "PG primary",
                ["Giao dịch + PII", "Partition · outbox · report"],
            ),
            (35, 410, 290, 95, "PG replica (profile)", ["WAL async · chưa tự promote"]),
            (
                405,
                230,
                290,
                125,
                "Mongo replica set",
                ["Metadata + projections", "3 member cùng host"],
            ),
            (
                770,
                230,
                290,
                125,
                "MinIO single-node",
                ["3 bucket private", "Ảnh/upload qua FastAPI"],
            ),
            (405, 410, 290, 95, "ETL worker", ["Poll 60s · upsert/retry/ACK"]),
        ],
        [
            (280, 100, 365, 100),
            (450, 155, 240, 225),
            (550, 155, 550, 225),
            (680, 120, 910, 225),
            (180, 355, 180, 405),
            (325, 305, 400, 452),
            (550, 410, 550, 360),
            (695, 452, 910, 360),
        ],
        "Một Docker host. Replication không phải sharding hoặc HA chịu mất host.",
    )
    svg_diagram(
        "07_erd_summary",
        [
            (25, 30, 265, 110, "user_core", ["PK user_id · roles qua user_role"]),
            (
                390,
                30,
                265,
                110,
                "product",
                ["PK product_id · seller_id", "price / quantity / status"],
            ),
            (780, 30, 265, 110, "category", ["PK category_id"]),
            (
                25,
                205,
                265,
                125,
                "user fragments",
                [
                    "profile / secret / wallet",
                    "PK/FK user_id → core",
                    "users_full: view chỉ đọc",
                ],
            ),
            (
                390,
                205,
                265,
                125,
                "order (HASH)",
                ["PK order_id · buyer_id", "contact / payment / voucher"],
            ),
            (
                780,
                205,
                265,
                125,
                "order_detail",
                ["FK order_id / product_id", "price_at_purchase / subtotal"],
            ),
            (
                390,
                400,
                265,
                100,
                "transaction",
                ["FK order_id", "UNIQUE reference_number"],
            ),
            (
                780,
                400,
                265,
                100,
                "product_image",
                ["FK product_id · URL / MIME / size"],
            ),
        ],
        [
            (290, 85, 385, 85),
            (775, 85, 660, 85),
            (158, 140, 158, 200),
            (290, 110, 400, 200),
            (775, 260, 660, 260),
            (520, 140, 850, 200),
            (520, 330, 520, 395),
            (1040, 140, 1040, 395),
        ],
        "ERD rút gọn; schema đầy đủ gồm vận hành ETL ở 07_erd_current.dbml.",
    )
    svg_diagram(
        "07_horizontal",
        [
            (35, 55, 270, 110, "order", ["HASH(order_id)", "PK đơn được giữ"]),
            (450, 30, 270, 80, "order_h0 / h1", ["modulus 4, remainder 0 / 1"]),
            (450, 145, 270, 80, "order_h2 / h3", ["modulus 4, remainder 2 / 3"]),
            (
                35,
                330,
                270,
                110,
                "system_log",
                ["RANGE(created_at)", "PK(log_id, created_at)"],
            ),
            (
                450,
                295,
                340,
                100,
                "system_log_YYYYMM",
                ["[đầu tháng, đầu tháng sau)", "Các khoảng không giao nhau"],
            ),
            (450, 430, 340, 80, "system_log_default", ["Nhận ngày ngoài lịch hiện có"]),
        ],
        [
            (305, 95, 445, 70),
            (305, 125, 445, 185),
            (305, 365, 445, 340),
            (305, 405, 445, 470),
        ],
        "Prune theo partition key; buyer/seller history không HASH-prune theo buyer/seller.",
    )
    svg_diagram(
        "07_vertical",
        [
            (35, 190, 260, 125, "user cũ", ["20 cột nguồn", "Tái tạo qua users_full"]),
            (
                450,
                25,
                550,
                100,
                "user_core",
                ["ID / username / email / trạng thái / thời gian"],
            ),
            (
                450,
                155,
                550,
                100,
                "user_profile",
                ["Tên / điện thoại / địa chỉ / avatar"],
            ),
            (
                450,
                285,
                550,
                100,
                "user_secret",
                ["Hash / random key / verification / reset tokens"],
            ),
            (
                450,
                415,
                550,
                100,
                "user_wallet",
                ["user_id / balance; đọc ghi tiền ở primary"],
            ),
        ],
        [
            (295, 225, 445, 75),
            (295, 245, 445, 205),
            (295, 270, 445, 335),
            (295, 290, 445, 465),
        ],
        "PK/FK user_id chung. Service tạo đủ các mảnh trong một transaction; view chỉ đọc.",
    )
    svg_diagram(
        "07_etl",
        [
            (
                20,
                60,
                240,
                115,
                "PG source",
                ["Transaction + triggers", "Key/op + provenance"],
            ),
            (
                325,
                60,
                240,
                115,
                "Outbox",
                ["pending / dead / done", "Không MAX(id) cursor"],
            ),
            (630, 60, 240, 115, "Worker", ["Exact VND · UTC", "Một advisory lock"]),
            (
                630,
                260,
                400,
                115,
                "Mongo projections",
                ["Partial upsert · majority write", "Giữ fields Mongo-owned → ACK"],
            ),
            (
                20,
                300,
                420,
                120,
                "PG aggregate / media cleanup",
                [
                    "report_daily_sales: gross theo ngày VN",
                    "Queue unlink, kiểm references, retry S3",
                ],
            ),
            (630, 445, 400, 75, "MinIO private", ["Byte object; không atomic với PG"]),
        ],
        [
            (260, 115, 320, 115),
            (565, 115, 625, 115),
            (750, 175, 750, 255),
            (630, 145, 440, 335),
            (830, 375, 830, 440),
        ],
        "At least once + idempotence. Poll 60 giây; quality count/field/media; lỗi giữ để retry.",
    )


def charts(pg, mongo):
    plt.rcParams.update(
        {
            "font.family": "DejaVu Sans",
            "font.size": 12,
            "axes.unicode_minus": False,
            "mathtext.fontset": "dejavusans",
        }
    )
    fig, ax = plt.subplots(figsize=(11.3, 5.8), layout="constrained")
    for i, key in enumerate(LABELS):
        before = pg["before"][key]["median_execution_ms"]
        after = pg["after"][key]["median_execution_ms"]
        ax.plot([before, after], [i, i], color="#C2CDD5", linewidth=2)
        ax.scatter(
            before,
            i,
            color="#DE763C",
            s=65,
            label="Trước" if i == 0 else None,
            zorder=3,
        )
        ax.scatter(
            after,
            i,
            color="#008F86",
            marker="D",
            s=55,
            label="Sau" if i == 0 else None,
            zorder=3,
        )
        ax.annotate(
            f"{before:g} → {after:g}",
            (max(before, after) * 1.15, i),
            fontsize=10,
            va="center",
        )
    ax.set(
        yticks=range(len(LABELS)),
        yticklabels=list(LABELS.values()),
        xscale="log",
        xlim=(0.025, 3500),
        xlabel="Execution Time (ms) — trung vị 3 lượt, trục log; thấp hơn là tốt hơn",
    )
    ax.invert_yaxis()
    ax.grid(axis="x", alpha=0.2)
    ax.legend(loc="lower right")
    ax.spines[["top", "right"]].set_visible(False)
    fig.savefig(ASSETS / "07_pg_benchmark.svg")
    plt.close(fig)
    cases = mongo["cases"]
    fig, ax = plt.subplots(figsize=(11.3, 5.5), layout="constrained")
    for i, (name, data) in enumerate(cases.items()):
        b = data["before"]["totalDocsExamined"]
        a = data["after"]["totalDocsExamined"]
        ax.plot([a, b], [i, i], color="#C2CDD5", linewidth=3)
        ax.scatter(b, i, color="#DE763C", s=100, label="Trước" if i == 0 else None)
        ax.scatter(
            a, i, color="#008F86", marker="D", s=90, label="Sau" if i == 0 else None
        )
        ax.text(b * 1.1, i, f"{b:,} → {a:,}", va="center", fontsize=13)
    ax.set(
        yticks=range(3),
        yticklabels=list(cases),
        xscale="log",
        xlim=(5, 60000),
        ylim=(2.6, -0.6),
        xlabel="Documents examined — 10.000 document/collection, một explain/pha",
    )
    ax.legend(loc="lower right")
    ax.grid(axis="x", alpha=0.2)
    ax.spines[["top", "right"]].set_visible(False)
    fig.savefig(ASSETS / "07_mongo_benchmark.svg")
    plt.close(fig)


def expand_report(pg, mongo):
    pgtable = "| Query | Trước (ms) | Sau (ms) |\n|---|---:|---:|\n" + "".join(
        f"| {label} | {pg['before'][key]['median_execution_ms']:.3f} | {pg['after'][key]['median_execution_ms']:.3f} |\n"
        for key, label in LABELS.items()
    )
    mtable = (
        "| Collection | Docs trước → sau | Returned | ms trước → sau |\n|---|---:|---:|---:|\n"
        + "".join(
            f"| {name} | {d['before']['totalDocsExamined']} → {d['after']['totalDocsExamined']} | {d['after']['nReturned']} | {d['before']['executionTimeMillis']} → {d['after']['executionTimeMillis']} |\n"
            for name, d in mongo["cases"].items()
        )
    )
    path = REPORT / "07_report.md"
    source = path.read_text()
    for name, value in [
        ("PG_BENCHMARK_TABLE", pgtable),
        ("MONGO_BENCHMARK_TABLE", mtable),
    ]:
        block = f"<!-- BEGIN {name} -->\n{value}<!-- END {name} -->"
        if "{{" + name + "}}" in source:
            source = source.replace("{{" + name + "}}", block)
        else:
            source = re.sub(
                f"<!-- BEGIN {name} -->.*?<!-- END {name} -->",
                lambda _, value=block: value,
                source,
                flags=re.DOTALL,
            )
    path.write_text(source)
    return source


def png_for(path):
    if path.suffix == ".svg":
        return io.BytesIO(cairosvg.svg2png(url=str(path), output_width=1800))
    return str(path)


def build_docx(source):
    doc = Document()
    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2.1)
    normal = doc.styles["Normal"]
    normal.font.name = "Noto Sans"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.16
    for style in ["Title", "Heading 1", "Heading 2", "Heading 3"]:
        doc.styles[style].font.name = "Noto Sans"
        doc.styles[style].font.color.rgb = RGBColor.from_string(NAVY)
    section.header.paragraphs[0].text = "OLDSHOP  /  BÁO CÁO KỸ THUẬT  /  30.10.2026"
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run("OldShop · ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    # Build markers belong to the Markdown source, not the exported document.
    tokens = MarkdownIt("default").parse(
        re.sub(r"<!--.*?-->", "", source, flags=re.DOTALL)
    )
    i = 0
    list_depth = 0
    while i < len(tokens):
        token = tokens[i]
        if token.type in ("bullet_list_open", "ordered_list_open"):
            list_depth += 1
        elif token.type in ("bullet_list_close", "ordered_list_close"):
            list_depth -= 1
        elif token.type == "heading_open":
            level = int(token.tag[1])
            text = tokens[i + 1].content
            if level == 1:
                doc.add_heading(text, 0)
            else:
                if text == "Tóm tắt":
                    doc.add_page_break()
                doc.add_heading(text, level - 1)
            i += 2
        elif token.type == "paragraph_open":
            inline = tokens[i + 1]
            images = [c for c in inline.children or [] if c.type == "image"]
            if images:
                for image in images:
                    path = REPORT / image.attrGet("src")
                    doc.add_picture(png_for(path), width=Cm(16.7))
                    p = doc.add_paragraph(image.content)
                    p.style = "Caption"
            else:
                p = doc.add_paragraph(style="List Bullet" if list_depth else None)
                for item in inline.children or []:
                    if item.type in ("text", "code_inline"):
                        run = p.add_run(item.content)
                        if item.type == "code_inline":
                            run.font.name = "Noto Sans Mono"
                            run.font.size = Pt(9)
                    elif item.type in ("softbreak", "hardbreak"):
                        p.add_run("\n")
                if not p.text:
                    p.add_run(inline.content)
            i += 2
        elif token.type == "fence":
            p = doc.add_paragraph()
            p.paragraph_format.keep_together = True
            r = p.add_run(token.content.rstrip())
            r.font.name = "Noto Sans Mono"
            r.font.size = Pt(8)
            p.paragraph_format.space_before = Pt(6)
        elif token.type == "table_open":
            rows = []
            row = []
            i += 1
            while tokens[i].type != "table_close":
                t = tokens[i]
                if t.type == "tr_open":
                    row = []
                elif t.type == "inline":
                    row.append(t.content.replace("**", "").replace("`", ""))
                elif t.type == "tr_close":
                    rows.append(row)
                i += 1
            if rows:
                table = doc.add_table(rows=1, cols=len(rows[0]))
                table.style = "Light Shading Accent 1"
                for cell, text in zip(table.rows[0].cells, rows[0]):
                    cell.text = text
                repeat = OxmlElement("w:tblHeader")
                table.rows[0]._tr.get_or_add_trPr().append(repeat)
                for row in rows[1:]:
                    for cell, text in zip(table.add_row().cells, row):
                        cell.text = text
                for row in table.rows:
                    for cell in row.cells:
                        for p in cell.paragraphs:
                            for run in p.runs:
                                run.font.size = Pt(8.5)
        i += 1
    doc.save(OUT / "oldshop_report.docx")


def slide_text(slide, text, x, y, w, h, size=23, color=NAVY, bold=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.word_wrap = True
    for i, line in enumerate(text.split("\n")):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.text = line
        p.font.name = "Noto Sans"
        p.font.size = PPTPt(size)
        p.font.bold = bold
        p.font.color.rgb = PPTColor.from_string(color)
        p.space_after = PPTPt(16)
    return box


def build_deck():
    deck = json.loads((REPORT / "07_slides.json").read_text())
    assert len(deck) == 18
    ppt = Presentation()
    ppt.slide_width = Inches(13.333)
    ppt.slide_height = Inches(7.5)
    md = ["# OldShop — 18 slide và lời dẫn\n"]
    for n, data in enumerate(deck, 1):
        slide = ppt.slides.add_slide(ppt.slide_layouts[6])
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = PPTColor.from_string(NAVY if n == 1 else "F6F8FA")
        color = "FFFFFF" if n == 1 else NAVY
        slide_text(
            slide,
            "OLDSHOP   /   DỮ LIỆU ĐA MÔ HÌNH",
            0.55,
            0.22,
            12,
            0.35,
            12,
            TEAL if n != 1 else "68DDC3",
            True,
        )
        slide_text(slide, data["title"], 0.55, 0.82, 12, 1.0, 34, color, True)
        if n == 1:
            slide_text(slide, data["subtitle"], 0.6, 2.0, 11, 1.3, 44, "FFFFFF", True)
            slide_text(slide, "\n".join(data["body"]), 0.65, 4.0, 11, 2, 24, "BBD5E1")
        elif data.get("visual"):
            img = ASSETS / (data["visual"] + ".svg")
            slide.shapes.add_picture(
                png_for(img),
                Inches(0.55),
                Inches(1.9),
                width=Inches(12.2),
                height=Inches(4.85),
            )
        elif (
            data.get("visual_optional")
            and (REPORT / "screenshots/04_product.png").exists()
        ):
            slide_text(
                slide,
                "\n".join("•  " + line for line in data.get("body", [])),
                0.7,
                2.05,
                5.2,
                4.7,
                22,
                color,
            )
            slide.shapes.add_picture(
                str(REPORT / "screenshots/04_product.png"),
                Inches(6.1),
                Inches(2.0),
                width=Inches(6.6),
            )
            slide_text(
                slide,
                "Trang sản phẩm thật trong project demo riêng",
                6.1,
                6.1,
                6.6,
                0.5,
                13,
                TEAL,
            )
        else:
            slide_text(
                slide,
                "\n".join("•  " + line for line in data.get("body", [])),
                0.7,
                2.05,
                11.9,
                4.6,
                25,
                color,
            )
        slide_text(
            slide,
            f"30/10/2026  ·  Một máy / Mongo7 local  ·  {n:02d}/18",
            0.65,
            7.05,
            12,
            0.25,
            10,
            "587082",
        )
        slide.notes_slide.notes_text_frame.text = data.get("notes", "")
        md += [
            f"\n## {n:02d}. {data['title']}\n",
            "\n".join("- " + b for b in data.get("body", [])),
            "\nLời dẫn: " + data.get("notes", "") + "\n",
        ]
        if data.get("visual"):
            md.append(f"![{data['title']}](assets/{data['visual']}.svg)\n")
        elif data.get("visual_optional"):
            md.append(
                "![Trang sản phẩm trong project demo](screenshots/04_product.png)\n"
            )
    ppt.save(OUT / "oldshop_slides.pptx")
    (REPORT / "07_slides.md").write_text("\n".join(md))


def convert_pdf():
    with tempfile.TemporaryDirectory(prefix="oldshop-lo-") as profile:
        for name in ["oldshop_report.docx", "oldshop_slides.pptx"]:
            subprocess.run(
                [
                    "libreoffice",
                    f"-env:UserInstallation=file://{profile}",
                    "--headless",
                    "--convert-to",
                    "pdf",
                    "--outdir",
                    str(OUT),
                    str(OUT / name),
                ],
                check=True,
                timeout=120,
            )
            assert (OUT / (Path(name).stem + ".pdf")).stat().st_size > 10000


def main():
    ASSETS.mkdir(exist_ok=True, parents=True)
    OUT.mkdir(exist_ok=True, parents=True)
    pg = json.loads((ROOT / "docs/benchmark/pg/summary.json").read_text())
    mongo = json.loads((ROOT / "docs/benchmark/mongo/summary.json").read_text())
    assert len(pg["before"]) >= 9 and mongo["version"] == "7.0.43"
    facts = {
        "report_date": "2026-10-30",
        "scope": "One-host MVP, migrations 001–006; no remote CI/Mongo8 claim",
        "backend_distinct_tests": 146,
        "frontend_tests": 15,
        "day6_failure_scenarios": 9,
        "pg_benchmark": pg,
        "mongo_benchmark": mongo,
        "sources": [
            "docs/evidence/day6_stack.json",
            "docs/evidence/day6_rollout.json",
            "docs/evidence/day6_api_contract.json",
            "docs/evidence/day6_backend_tests.txt",
            "docs/evidence/day6_final_tests.txt",
            "docs/evidence/day6_status_regression.txt",
        ],
    }
    (REPORT / "07_facts.json").write_text(
        json.dumps(facts, ensure_ascii=False, indent=2) + "\n"
    )
    diagrams()
    charts(pg, mongo)
    # Matplotlib SVG paths contain trailing spaces; normalize source without changing drawing.
    for path in ASSETS.glob("*.svg"):
        path.write_text(
            "\n".join(line.rstrip() for line in path.read_text().splitlines()).rstrip()
            + "\n"
        )
    source = expand_report(pg, mongo)
    build_docx(source)
    build_deck()
    convert_pdf()
    print("PASS: DOCX/PDF report, 18-slide PPTX/PDF, SVG diagrams/charts.")


if __name__ == "__main__":
    main()
